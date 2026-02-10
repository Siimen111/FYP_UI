import os

from docx import Document


def write_docx(snapshot: dict, path: str) -> None:
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)

    doc = Document()
    title = snapshot.get("title") or snapshot.get("doc_id") or "Merged Document"
    doc.add_heading(title, level=1)

    clauses = snapshot.get("clauses") or []
    for clause in clauses:
        cid = clause.get("cid") or ""
        heading = clause.get("heading") or ""
        if cid and heading:
            label = f"{cid} - {heading}"
        else:
            label = heading or cid or "Clause"
        doc.add_heading(label, level=2)

        text = clause.get("text") or ""
        if text:
            for line in text.split("\n"):
                doc.add_paragraph(line)
        else:
            doc.add_paragraph("")

    doc.save(path)
